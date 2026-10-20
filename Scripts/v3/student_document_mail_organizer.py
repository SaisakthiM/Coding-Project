#!/usr/bin/env python3
"""Fetch email attachments and organize student documents by student and document type.

Dependencies: pip install pypdf pillow pytesseract openpyxl pymupdf
Also install the Tesseract OCR application and ensure `tesseract` is on PATH.

Set IMAP_HOST, IMAP_USER, IMAP_PASSWORD as environment variables. For Gmail, use an
app password when required. This script uses read-only IMAP mailbox selection.
"""
import argparse, csv, email, imaplib, os, re, shutil, sys, zipfile
from email.header import decode_header
from pathlib import Path

try:
    from pypdf import PdfReader, PdfWriter
except ImportError:
    PdfReader = PdfWriter = None
try:
    from PIL import Image
except ImportError:
    Image = None
try:
    import pytesseract
except ImportError:
    pytesseract = None
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

PATTERNS = {
 "10th Marksheet": [r"\b10\s*(?:TH|STD|STANDARD|CLASS)\b", r"\bSSLC\b", r"\bSECONDARY SCHOOL LEAVING CERTIFICATE\b"],
 "11th Marksheet": [r"\b11\s*(?:TH|STD|STANDARD|CLASS)\b", r"\bXI\s*(?:STD|STANDARD)?\b", r"\bHIGHER SECONDARY FIRST YEAR\b"],
 "12th Marksheet": [r"\b12\s*(?:TH|STD|STANDARD|CLASS)\b", r"\bHSC\b", r"\bXII\s*(?:STD|STANDARD)?\b", r"\bHIGHER SECONDARY CERTIFICATE\b", r"\bMARK STATEMENT\b"],
 "TC": [r"\bTRANSFER CERTIFICATE\b", r"\bT\.?\s*C\.?\b", r"\bSCHOOL LEAVING CERTIFICATE\b"],
 "Community Certificate": [r"\bCOMMUNITY CERTIFICATE\b", r"\bCASTE CERTIFICATE\b"],
 "Income Certificate": [r"\bINCOME CERTIFICATE\b", r"\bANNUAL INCOME\b"],
 "Nativity Certificate": [r"\bNATIVITY CERTIFICATE\b", r"\bNATIVE OF TAMIL NADU\b"],
 "Aadhar Card": [r"\bAADHAAR\b", r"\bAADHAR\b", r"\bUIDAI\b"],
 "PAN Card": [r"\bPERMANENT ACCOUNT NUMBER\b", r"\bINCOME TAX DEPARTMENT\b"],
 "First Graduate Certificate": [r"\bFIRST GRADUATE\b", r"\bFIRST GRADUATE CERTIFICATE\b"],
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

def dec(value):
    if not value: return ""
    out = []
    for part, enc in decode_header(value):
        out.append(part.decode(enc or "utf-8", errors="replace") if isinstance(part, bytes) else part)
    return "".join(out)

def safe(s):
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(s or "Unknown"))
    return re.sub(r"\s+", " ", s).strip(" ._-")[:120] or "Unknown"

def unique(path):
    if not path.exists(): return path
    n, stem, ext = 2, path.stem, path.suffix
    while (path.parent / f"{stem}_{n}{ext}").exists(): n += 1
    return path.parent / f"{stem}_{n}{ext}"

def norm(s): return re.sub(r"[^A-Z0-9]+", " ", str(s).upper()).strip()

def load_names(xlsx, sheet):
    if not xlsx: return []
    if load_workbook is None: raise RuntimeError("Install openpyxl: pip install openpyxl")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [norm(x) for x in next(rows, ())]
    idx = next((headers.index(h) for h in ("STUDENT NAME", "NAME", "CANDIDATE NAME") if h in headers), None)
    if idx is None: raise ValueError(f"No student name column found. Headers: {headers}")
    result = []
    for row in rows:
        if idx < len(row) and row[idx]:
            name = re.sub(r"\s+", " ", str(row[idx])).strip()
            if name and name not in result: result.append(name)
    wb.close()
    return result

def find_student(text, names):
    if not names: return None
    hay = f" {norm(text)} "
    matches = []
    for name in names:
        n = norm(name)
        if n and f" {n} " in hay: matches.append(name); continue
        toks = n.split()
        if len(toks) >= 2 and all(re.search(rf"\b{re.escape(t)}\b", hay) for t in toks):
            matches.append(name)
    return matches[0] if len(matches) == 1 else None

def classify(text):
    text = re.sub(r"\s+", " ", (text or "").upper())
    hits = []
    for doc, patterns in PATTERNS.items():
        score = sum(bool(re.search(p, text, re.I)) for p in patterns)
        if score: hits.append((score, doc))
    hits.sort(reverse=True)
    if not hits or (len(hits) > 1 and hits[0][0] == hits[1][0]): return None
    return hits[0][1]

def image_text(path):
    if not Image or not pytesseract: return ""
    try:
        with Image.open(path) as im: return pytesseract.image_to_string(im)
    except Exception: return ""

def pdf_texts(path):
    if not PdfReader: raise RuntimeError("Install pypdf: pip install pypdf")
    reader = PdfReader(str(path))
    texts = [p.extract_text() or "" for p in reader.pages]
    if any(len(t.strip()) < 20 for t in texts) and pytesseract:
        try:
            import fitz
            doc = fitz.open(str(path))
            for i, text in enumerate(texts):
                if len(text.strip()) < 20:
                    pix = doc[i].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    import io
                    texts[i] = pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))))
            doc.close()
        except Exception:
            pass
    return reader, texts

def process_pdf(path, outroot, hint, names, report):
    reader, texts = pdf_texts(path)
    groups = []
    for i, text in enumerate(texts):
        doc_type = classify(text)
        student = find_student(text, names) or hint
        if groups and groups[-1][0] == student and groups[-1][1] == doc_type and doc_type:
            groups[-1][2].append(i)
        else:
            groups.append([student, doc_type, [i]])
    for student, doc_type, pages in groups:
        target = outroot / (safe(student) if doc_type else f"_REVIEW/{safe(student or hint)}")
        target.mkdir(parents=True, exist_ok=True)
        label = safe(doc_type or "REVIEW")
        page_tag = f"_pages_{pages[0]+1}-{pages[-1]+1}" if len(pages) > 1 else f"_page_{pages[0]+1}"
        dest = unique(target / f"{safe(student or hint)}_{label}{page_tag}.pdf")
        writer = PdfWriter()
        for i in pages: writer.add_page(reader.pages[i])
        with dest.open("wb") as f: writer.write(f)
        report.append({"source": str(path), "student": student or hint or "", "document": doc_type or "REVIEW",
                       "pages": ",".join(str(i+1) for i in pages), "output": str(dest),
                       "status": "classified" if doc_type else "review"})

def process_file(path, outroot, hint, names, report):
    ext = path.suffix.lower()
    if ext == ".pdf":
        process_pdf(path, outroot, hint, names, report); return
    if ext in IMAGE_EXTS:
        text = image_text(path)
        doc_type = classify(text)
        student = find_student(text, names) or hint
        target = outroot / (safe(student) if doc_type else f"_REVIEW/{safe(student or hint)}")
        target.mkdir(parents=True, exist_ok=True)
        dest = unique(target / f"{safe(student or hint)}_{safe(doc_type or 'REVIEW')}{ext}")
        shutil.copy2(path, dest)
        report.append({"source": str(path), "student": student or hint or "", "document": doc_type or "REVIEW",
                       "pages": "", "output": str(dest), "status": "classified" if doc_type else "review"})
        return
    target = outroot / "_REVIEW" / safe(hint)
    target.mkdir(parents=True, exist_ok=True)
    dest = unique(target / path.name); shutil.copy2(path, dest)
    report.append({"source": str(path), "student": hint or "", "document": "UNSUPPORTED", "pages": "",
                   "output": str(dest), "status": "review"})

def extract_zip(path):
    root = path.parent / (path.stem + "_extracted")
    root.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            member = Path(info.filename)
            if member.is_absolute() or ".." in member.parts: continue
            if info.is_dir(): continue
            dest = root / member
            dest.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, dest.open("wb") as dst:
                shutil.copyfileobj(src, dst)
    return root

def mail_fetch(args):
    host = args.imap_host or os.getenv("IMAP_HOST")
    user = args.imap_user or os.getenv("IMAP_USER")
    password = args.imap_password or os.getenv("IMAP_PASSWORD")
    if not all([host, user, password]):
        raise ValueError("Set IMAP_HOST, IMAP_USER, IMAP_PASSWORD or use the matching --imap-* options.")
    staging = args.out / "_mail_downloads"; staging.mkdir(parents=True, exist_ok=True)
    conn = imaplib.IMAP4_SSL(host, args.imap_port)
    conn.login(user, password)
    status, _ = conn.select(args.mailbox, readonly=True)
    if status != "OK": raise RuntimeError(f"Cannot open mailbox {args.mailbox}")
    status, data = conn.search(None, "UNSEEN" if args.unseen_only else "ALL")
    if status != "OK": raise RuntimeError("IMAP search failed")
    ids = data[0].split()
    if args.limit: ids = ids[-args.limit:]
    messages = []
    for mid in ids:
        status, payload = conn.fetch(mid, "(RFC822)")
        if status != "OK" or not payload or not payload[0]: continue
        msg = email.message_from_bytes(payload[0][1])
        subject = dec(msg.get("Subject"))
        bodies, attachments = [], []
        for part in msg.walk():
            if part.is_multipart(): continue
            disp = part.get("Content-Disposition", "").lower()
            filename = dec(part.get_filename())
            if "attachment" in disp or filename:
                if not filename: filename = f"attachment_{mid.decode()}_{len(attachments)+1}.bin"
                filename = safe(filename)
                data_bytes = part.get_payload(decode=True)
                if data_bytes is not None:
                    saved = unique(staging / f"{mid.decode()}_{filename}")
                    saved.write_bytes(data_bytes); attachments.append(saved)
            elif part.get_content_type() == "text/plain":
                try: bodies.append(part.get_payload(decode=True).decode(part.get_content_charset() or "utf-8", "replace"))
                except Exception: pass
        messages.append((mid.decode(), subject, "\n".join(bodies), attachments))
    conn.logout()
    return staging, messages

def main():
    ap = argparse.ArgumentParser(description="Fetch student document emails and organize attachments.")
    ap.add_argument("--out", type=Path, default=Path("student_documents"))
    ap.add_argument("--excel", type=Path, help="Roster Excel file with STUDENT NAME column (recommended)")
    ap.add_argument("--sheet", help="Roster sheet name")
    ap.add_argument("--input-dir", type=Path, help="Process existing files instead of fetching email")
    ap.add_argument("--imap-host", help="IMAP hostname; alternatively set IMAP_HOST")
    ap.add_argument("--imap-port", type=int, default=993)
    ap.add_argument("--imap-user", help="Email login; alternatively set IMAP_USER")
    ap.add_argument("--imap-password", help="Email app password; alternatively set IMAP_PASSWORD")
    ap.add_argument("--mailbox", default="INBOX")
    ap.add_argument("--unseen-only", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help="Most recent N emails; 0 means all")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    names = load_names(args.excel, args.sheet)
    report = []
    if args.input_dir:
        files = list(args.input_dir.rglob("*"))
        for z in [p for p in files if p.is_file() and p.suffix.lower() == ".zip"]:
            extract_zip(z)
        for p in args.input_dir.rglob("*"):
            if p.is_file() and p.suffix.lower() != ".zip" and "_extracted" not in str(p):
                process_file(p, args.out, None, names, report)
    else:
        staging, messages = mail_fetch(args)
        for mid, subject, body, attachments in messages:
            hint = find_student(subject + "\n" + body, names)
            for attachment in attachments:
                message_dir = staging / f"message_{mid}"; message_dir.mkdir(parents=True, exist_ok=True)
                local = unique(message_dir / attachment.name); shutil.copy2(attachment, local)
                if local.suffix.lower() == ".zip":
                    root = extract_zip(local)
                    for p in root.rglob("*"):
                        if p.is_file(): process_file(p, args.out, hint, names, report)
                else:
                    process_file(local, args.out, hint, names, report)
    with (args.out / "processing_report.csv").open("w", newline="", encoding="utf-8-sig") as f:
        fields = ["source", "student", "document", "pages", "output", "status"]
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(report)
    print(f"Processed {len(report)} document item(s). Output: {args.out.resolve()}")
    print(f"Review uncertain items in {args.out / '_REVIEW'}; report: {args.out / 'processing_report.csv'}")
    return 0

if __name__ == "__main__":
    try: raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr); raise SystemExit(1)
